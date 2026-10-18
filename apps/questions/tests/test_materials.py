"""Learning materials list, download, and student library."""

from __future__ import annotations

from unittest.mock import patch

import pytest
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.ai.ingest import ingest_learning_upload
from apps.ai.models import LearningDocument
from apps.ai.source_extract import SourceMaterialError
from apps.reviews.exam_setup_services import sync_program_subject_timers
from apps.users.assignment_services import get_or_create_catalog_course
from conftest import make_bsed_student


def _make_ready_doc(*, course, subject, user, title="Module overview", archived=False):
    doc = LearningDocument.objects.create(
        course_id=course.pk,
        subject=subject,
        created_by=user,
        title=title,
        original_name=f"{title}.txt",
        material_type=LearningDocument.MaterialType.MODULE,
        status=LearningDocument.Status.READY,
        page_count=1,
        chunk_count=1,
        is_archived=archived,
    )
    doc.file.save(f"{title}.txt", ContentFile(b"Sample module content."), save=True)
    return doc


@pytest.mark.django_db
class TestFacultyMaterialsUI:
    def test_course_list_keeps_course_chrome(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(professor)
        response = client.get(
            reverse("analytics_professor:material_list", kwargs={"course_pk": course.pk})
        )
        assert response.status_code == 200
        content = response.content.decode()
        assert "Content management" in content
        assert "materials-library" in content
        # Course sub-navigation was dropped; the sidebar owns those links now.
        assert "course-nav-link" not in content
        assert ">Download</a>" in content
        assert ">Edit</a>" in content
        # Course-scoped page: no course picker, supported-types hint present.
        assert "materials-hub-upload-row" in content
        assert "Supported: TXT, MD, CSV, PDF, DOCX." in content
        assert "Choose file" in content
        assert "<select" not in content.split("materials-library")[0]

    def test_hub_with_subject_shows_download_and_edit(
        self, client, professor, subject, bsed_program
    ):
        subject.program = bsed_program
        subject.save(update_fields=["program"])
        course = get_or_create_catalog_course(professor, subject)
        _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(professor)
        response = client.get(
            reverse("analytics_professor:materials_hub") + f"?subject={subject.pk}"
        )
        content = response.content.decode()
        assert response.status_code == 200
        assert "materials-library" in content
        assert ">Download</a>" in content
        assert ">Edit</a>" in content
        assert "Archive" in content
        assert "Content Management" in content

    def test_download_returns_attachment(self, client, professor, subject):
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(professor)
        response = client.get(
            reverse(
                "analytics_professor:material_download",
                kwargs={"course_pk": course.pk, "pk": doc.pk},
            )
        )
        assert response.status_code == 200
        assert response["Content-Disposition"].startswith("attachment;")
        assert b"Sample module content." in b"".join(response.streaming_content)

    def test_archive_from_list_hides_from_students(
        self, client, professor, student, subject, bsed_program
    ):
        make_bsed_student(student, subject=subject, bsed_program=bsed_program)
        sync_program_subject_timers(
            bsed_program,
            subject_ids=[subject.pk],
            timers_by_subject_id={subject.pk: 30},
            default_seconds=30,
        )
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(professor)
        archive = client.post(
            reverse(
                "analytics_professor:material_detail",
                kwargs={"course_pk": course.pk, "pk": doc.pk},
            ),
            {"action": "archive"},
        )
        assert archive.status_code == 302
        doc.refresh_from_db()
        assert doc.is_archived is True

        client.force_login(student)
        listing = client.get(
            reverse(
                "analytics_student:materials_list",
                kwargs={"subject_pk": subject.pk},
            )
        )
        content = listing.content.decode()
        assert doc.display_title not in content
        download = client.get(
            reverse(
                "analytics_student:materials_download",
                kwargs={"subject_pk": subject.pk, "pk": doc.pk},
            )
        )
        assert download.status_code == 404


@pytest.fixture(autouse=True)
def _no_ai(settings):
    """Keep the relevance gate on its deterministic keyword path."""
    settings.AI_ENABLED = False


RELATED_UPLOAD = (
    b"Trigonometry module. This lesson covers the unit circle, sine, cosine and "
    b"tangent ratios, radians versus degrees, and solving right triangles using "
    b"trigonometric identities throughout worked examples."
)

UNRELATED_UPLOAD = (
    b"Photosynthesis and cellular respiration in plants. Chloroplasts capture "
    b"sunlight, chlorophyll pigments absorb energy, and glucose is produced "
    b"during the Calvin cycle in the stroma of every green leaf."
)


class _AlwaysUnrelated:
    """Patch the relevance gate so tests do not depend on keyword scoring."""

    @staticmethod
    def __call__(sample, subject, topics):
        return {
            "related": False,
            "matched_topic_id": None,
            "subject_score": 0,
            "topic_score": 0,
            "reason": "This module does not look related to the selected subject.",
        }


class _AlwaysRelated:
    @staticmethod
    def __call__(sample, subject, topics):
        return {
            "related": True,
            "matched_topic_id": None,
            "subject_score": 5,
            "topic_score": 5,
            "reason": "",
        }


@pytest.mark.django_db
class TestMaterialUploadRelevanceGate:
    """Uploads that do not match the course subject are rejected."""

    def _upload_url(self, course):
        return reverse(
            "analytics_professor:material_upload", kwargs={"course_pk": course.pk}
        )

    def _post(self, client, course, subject, payload):
        return client.post(
            self._upload_url(course),
            {"subject_id": subject.pk, **payload},
            follow=True,
        )

    def test_unrelated_upload_is_blocked_and_creates_nothing(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)

        with patch(
            "apps.questions.views_materials.assess_subject_relevance",
            _AlwaysUnrelated(),
        ):
            response = self._post(
                client,
                course,
                subject,
                {
                    "source_file": SimpleUploadedFile(
                        "plant-science.txt", UNRELATED_UPLOAD, content_type="text/plain"
                    )
                },
            )

        assert response.status_code == 200
        content = response.content.decode()
        assert "does not look related" in content
        assert LearningDocument.objects.filter(course_id=course.pk).count() == 0

    def test_related_upload_is_accepted(self, client, professor, subject):
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)

        with patch(
            "apps.questions.views_materials.assess_subject_relevance",
            _AlwaysRelated(),
        ):
            self._post(
                client,
                course,
                subject,
                {
                    "source_file": SimpleUploadedFile(
                        "trig.txt", RELATED_UPLOAD, content_type="text/plain"
                    )
                },
            )

        assert LearningDocument.objects.filter(course_id=course.pk).count() == 1

    def test_relevance_is_assessed_before_anything_is_written(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)
        seen = {}

        def _capture(sample, subj, topics):
            seen["sample"] = sample
            seen["subject"] = subj
            return {
                "related": False,
                "matched_topic_id": None,
                "reason": "nope",
                "subject_score": 0,
                "topic_score": 0,
            }

        with patch(
            "apps.questions.views_materials.assess_subject_relevance", _capture
        ):
            self._post(
                client,
                course,
                subject,
                {
                    "source_file": SimpleUploadedFile(
                        "plant.txt", UNRELATED_UPLOAD, content_type="text/plain"
                    )
                },
            )

        assert seen["subject"] == subject
        assert "Photosynthesis" in seen["sample"]

    def test_unreadable_upload_is_allowed_with_a_warning(
        self, client, professor, subject
    ):
        """No readable text means we cannot judge; keep the file downloadable."""
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)
        blob = SimpleUploadedFile(
            "scan.pdf", b"%PDF-1.4 tiny", content_type="application/pdf"
        )

        with patch(
            "apps.questions.views_materials.assess_subject_relevance"
        ) as mock_relevance:
            response = self._post(
                client, course, subject, {"source_file": blob}
            )

        mock_relevance.assert_not_called()
        assert response.status_code == 200
        assert LearningDocument.objects.filter(course_id=course.pk).count() == 1

    def test_unsupported_extension_is_rejected(
        self, client, professor, subject
    ):
        """Only TXT, MD, CSV, PDF and DOCX may be uploaded."""
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)
        blob = SimpleUploadedFile(
            "slides.pptx", b"PK\x03\x04not-really", content_type=(
                "application/vnd.openxmlformats-officedocument"
            )
        )

        with patch(
            "apps.questions.views_materials.assess_subject_relevance",
            _AlwaysRelated(),
        ):
            response = self._post(
                client, course, subject, {"source_file": blob}
            )

        assert response.status_code == 200
        assert LearningDocument.objects.filter(course_id=course.pk).count() == 0
        assert "not a supported format" in response.content.decode()

    def test_allowed_extensions_are_exactly_the_supported_set(self):
        from apps.ai.ingest import allowed_upload_suffixes, is_allowed_upload

        assert set(allowed_upload_suffixes()) == {
            ".txt",
            ".md",
            ".csv",
            ".pdf",
            ".docx",
        }
        assert is_allowed_upload("Lesson.PDF") is True
        assert is_allowed_upload("notes.docx") is True
        assert is_allowed_upload("deck.pptx") is False
        assert is_allowed_upload("sheet.xlsx") is False
        assert is_allowed_upload("photo.png") is False
        assert is_allowed_upload("") is False

    def test_blank_subject_falls_back_to_the_course_subject(
        self, client, professor, subject
    ):
        """The course-scoped form omits subject_id; the gate must still run."""
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)
        seen = {}

        def _capture(sample, subj, topics):
            seen["subject"] = subj
            return {
                "related": True,
                "matched_topic_id": None,
                "reason": "",
                "subject_score": 5,
                "topic_score": 5,
            }

        with patch(
            "apps.questions.views_materials.assess_subject_relevance", _capture
        ):
            response = client.post(
                self._upload_url(course),
                {
                    "source_file": SimpleUploadedFile(
                        "trig.txt", RELATED_UPLOAD, content_type="text/plain"
                    )
                },
                follow=True,
            )

        assert response.status_code == 200
        assert seen["subject"] == subject
        doc = LearningDocument.objects.get(course_id=course.pk)
        assert doc.subject == subject


@pytest.mark.django_db
class TestMaterialUploadFormTrimmed:
    """The upload surfaces dropped the Type field and the section heading."""

    def test_hub_has_helper_text_and_no_type_field(self, client, professor, subject):
        client.force_login(professor)
        content = client.get(
            reverse("analytics_professor:materials_hub") + f"?subject={subject.pk}"
        ).content.decode()

        assert "Supported: TXT, MD, CSV, PDF, DOCX." in content
        assert "Upload material" not in content
        assert 'name="material_type"' not in content
        assert 'id="material-type"' not in content

    def test_course_list_has_helper_text_and_no_type_field(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        client.force_login(professor)
        content = client.get(
            reverse(
                "analytics_professor:material_list",
                kwargs={"course_pk": course.pk},
            )
        ).content.decode()

        assert "Supported: TXT, MD, CSV, PDF, DOCX." in content
        assert "Upload material" not in content
        assert 'name="material_type"' not in content


@pytest.mark.django_db
class TestMaterialContentEditing:
    """Faculty can edit a material's content in place, Google-Docs style."""

    @pytest.fixture(autouse=True)
    def _isolated_media(self, settings, tmp_path):
        """Media is not rolled back with the transaction.

        Sharing MEDIA_ROOT across runs leaves thousands of suffixed leftovers,
        which makes Django pick random names and turns these assertions
        order-dependent.
        """
        settings.MEDIA_ROOT = str(tmp_path / "media")

    def _detail_url(self, course, doc):
        return reverse(
            "analytics_professor:material_detail",
            kwargs={"course_pk": course.pk, "pk": doc.pk},
        )

    def test_detail_page_renders_document_header(
        self, client, professor, subject
    ):
        """Document title, actions, meta pills and the overflow menu."""
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)
        client.force_login(professor)
        content = client.get(self._detail_url(course, doc)).content.decode()

        # Back link, breadcrumb, and course name removed per new design
        assert "materials-doc__crumbs" not in content
        assert "overview-back-link" not in content
        assert "Back to content management" not in content
        assert "All changes saved" in content
        assert "Generate questions" in content
        assert "data-dropdown-toggle" in content
        # _make_ready_doc stores a .txt, so the menu reads "Download txt".
        assert "Download txt" in content
        assert "Rename file" in content
        assert "Move to course" in content
        assert "Archive material" in content
        # Course picker is a pill select in the meta row.
        assert 'name="subject_id"' in content
        assert 'data-doc-rename-row' in content
        assert "materials-detail.js" in content

    def test_detail_page_renders_editor_with_current_text(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)
        doc.source_text = "Editable body text for the module."
        doc.save(update_fields=["source_text"])

        client.force_login(professor)
        content = client.get(self._detail_url(course, doc)).content.decode()

        assert "materials-doc__editor" in content
        assert 'data-editor-surface' in content
        assert "Editable body text for the module." in content
        assert "Text preview" not in content

    def test_saving_content_rebuilds_chunks_for_retrieval(
        self, client, professor, subject
    ):
        from apps.ai.models import LearningChunk

        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)
        doc.source_text = "Old content that should be replaced entirely."
        doc.save(update_fields=["source_text"])

        client.force_login(professor)
        response = client.post(
            self._detail_url(course, doc),
            {
                "action": "save_content",
                "content_html": "<h1>Heading</h1><p>Brand new edited content.</p>",
            },
        )

        assert response.status_code == 302
        doc.refresh_from_db()
        # source_text is the plain-text projection the editor HTML flattens to,
        # and it is what AI retrieval reads.
        assert doc.source_text == "# Heading\n\nBrand new edited content."
        assert "<h1>Heading</h1>" in doc.content_html
        assert doc.chunk_count == 1
        chunk_text = " ".join(
            LearningChunk.objects.filter(document=doc).values_list("text", flat=True)
        )
        assert "Brand new edited content." in chunk_text
        assert "Old content" not in chunk_text

    def test_saving_content_rewrites_plain_text_download(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(professor)
        client.post(
            self._detail_url(course, doc),
            {"action": "save_content", "content_html": "<p>Edited download body.</p>"},
        )

        download = client.get(
            reverse(
                "analytics_professor:material_download",
                kwargs={"course_pk": course.pk, "pk": doc.pk},
            )
        )
        assert b"Edited download body." in b"".join(download.streaming_content)

    def test_docx_download_is_regenerated_from_editor_html(
        self, client, professor, subject
    ):
        """Round-trip: the exported .docx mirrors what the editor showed."""
        import io

        import docx as docx_lib

        course = get_or_create_catalog_course(professor, subject)
        doc = LearningDocument.objects.create(
            course_id=course.pk,
            subject=subject,
            created_by=professor,
            title="Trig notes",
            original_name="trig-notes.docx",
            status=LearningDocument.Status.READY,
        )
        doc.file.save("trig-notes.docx", ContentFile(b"PK\x03\x04stub"), save=True)

        client.force_login(professor)
        client.post(
            self._detail_url(course, doc),
            {
                "action": "save_content",
                "content_html": (
                    "<h1>Unit Circle</h1>"
                    "<p>The <strong>sine</strong> ratio matters.</p>"
                    "<ul><li>First</li><li>Second</li></ul>"
                ),
            },
        )

        download = client.get(
            reverse(
                "analytics_professor:material_download",
                kwargs={"course_pk": course.pk, "pk": doc.pk},
            )
        )
        payload = b"".join(download.streaming_content)
        assert payload[:2] == b"PK"  # a real .docx zip container

        exported = docx_lib.Document(io.BytesIO(payload))
        styles = {p.style.name for p in exported.paragraphs}
        texts = [p.text for p in exported.paragraphs]
        assert "Heading 1" in styles
        assert "List Bullet" in styles
        assert "Unit Circle" in texts
        assert any("sine" in t for t in texts)
        assert any(r.bold for p in exported.paragraphs for r in p.runs)

    def test_saved_html_is_sanitized_server_side(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(professor)
        client.post(
            self._detail_url(course, doc),
            {
                "action": "save_content",
                "content_html": (
                    "<p onclick='steal()'>Safe text</p>"
                    "<script>alert(1)</script>"
                    "<img src='javascript:alert(1)'>"
                ),
            },
        )

        doc.refresh_from_db()
        assert "<script" not in doc.content_html
        assert "onclick" not in doc.content_html
        assert "javascript:" not in doc.content_html
        assert "Safe text" in doc.content_html

    def test_ingest_converts_docx_structure_to_html(self, professor, course):
        import io

        import docx as docx_lib

        source = docx_lib.Document()
        source.add_heading("Unit Circle", level=1)
        para = source.add_paragraph("The ")
        para.add_run("sine").bold = True
        para.add_run(" ratio.")
        source.add_paragraph("Bullet item", style="List Bullet")
        buffer = io.BytesIO()
        source.save(buffer)
        buffer.seek(0)

        upload = SimpleUploadedFile(
            "trig.docx", buffer.read(), content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )
        doc = ingest_learning_upload(
            uploaded_file=upload, course_id=course.pk, user=professor
        )

        assert "<h1>Unit Circle</h1>" in doc.content_html
        assert "<strong>sine</strong>" in doc.content_html
        assert "<ul><li>Bullet item</li></ul>" in doc.content_html
        # Plain-text projection is what AI retrieval reads.
        assert "# Unit Circle" in doc.source_text
        assert "sine" in doc.source_text

    def test_rewritten_file_path_is_not_double_nested(
        self, client, professor, subject
    ):
        course = get_or_create_catalog_course(professor, subject)
        doc = LearningDocument.objects.create(
            course_id=course.pk,
            subject=subject,
            created_by=professor,
            title="Path check",
            original_name="path-check-unique.txt",
            status=LearningDocument.Status.READY,
        )
        doc.file.save("path-check-unique.txt", ContentFile(b"before"), save=True)
        original_path = doc.file.name

        client.force_login(professor)
        client.post(
            self._detail_url(course, doc),
            {"action": "save_content", "content_html": "<p>Path check body.</p>"},
        )

        doc.refresh_from_db()
        parts = doc.file.name.split("/")
        # upload_to is "learning_modules/<year>/<month>/" and must appear once;
        # the old name must not be nested underneath itself.
        assert parts.count("learning_modules") == 1
        assert doc.file.name == original_path
        assert parts[-1] == "path-check-unique.txt"

    def test_binary_download_is_kept_but_text_still_updates(
        self, client, professor, subject
    ):
        from apps.ai.models import LearningChunk

        course = get_or_create_catalog_course(professor, subject)
        doc = LearningDocument.objects.create(
            course_id=course.pk,
            subject=subject,
            created_by=professor,
            title="Scanned notes",
            original_name="scanned.pdf",
            status=LearningDocument.Status.READY,
        )
        doc.file.save("scanned.pdf", ContentFile(b"%PDF-1.4 fake"), save=True)
        original_bytes = b"%PDF-1.4 fake"

        client.force_login(professor)
        response = client.post(
            self._detail_url(course, doc),
            {
                "action": "save_content",
                "content_html": "<p>Typed up by hand.</p>",
            },
        )

        assert response.status_code == 302
        doc.refresh_from_db()
        assert doc.source_text == "Typed up by hand."
        # Binary original stays downloadable for students.
        doc.file.open("rb")
        try:
            assert doc.file.read() == original_bytes
        finally:
            doc.file.close()
        assert LearningChunk.objects.filter(document=doc).exists()

    def test_blank_content_is_rejected(self, client, professor, subject):
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)
        doc.source_text = "Keep me."
        doc.save(update_fields=["source_text"])

        client.force_login(professor)
        response = client.post(
            self._detail_url(course, doc),
            {"action": "save_content", "content_html": "   "},
            follow=True,
        )

        doc.refresh_from_db()
        assert doc.source_text == "Keep me."
        assert b"Add some content before saving." in response.content

    def test_ingest_populates_source_text(self, professor, course):
        upload = SimpleUploadedFile(
            "module.txt",
            b"Linear equations in one variable. Solve for x carefully.",
            content_type="text/plain",
        )
        doc = ingest_learning_upload(
            uploaded_file=upload, course_id=course.pk, user=professor
        )
        assert "Linear equations" in doc.source_text
        assert doc.status == LearningDocument.Status.READY


@pytest.mark.django_db
class TestDocumentPlainTextFallback:
    """Documents predating source_text still round-trip their chunks."""

    def test_reassembles_chunks_without_duplicating_overlap(self, professor, course):
        from apps.ai.material_services import document_plain_text
        from apps.ai.source_extract import chunk_pages

        body = ("word " * 400).strip()  # long enough to split across chunks
        doc = LearningDocument.objects.create(
            course_id=course.pk,
            created_by=professor,
            original_name="x.txt",
            status=LearningDocument.Status.READY,
            source_text="",
        )
        for item in chunk_pages([{"page": 1, "text": body}]):
            doc.chunks.create(
                order=item["order"],
                page_start=item["page_start"],
                page_end=item["page_end"],
                text=item["text"],
            )

        assert doc.chunks.count() > 1
        assert document_plain_text(doc) == body


@pytest.mark.django_db
class TestStudentMaterials:
    def test_hub_and_subject_library_with_download(
        self, client, student, professor, subject, bsed_program
    ):
        make_bsed_student(student, subject=subject, bsed_program=bsed_program)
        sync_program_subject_timers(
            bsed_program,
            subject_ids=[subject.pk],
            timers_by_subject_id={subject.pk: 30},
            default_seconds=30,
        )
        course = get_or_create_catalog_course(professor, subject)
        doc = _make_ready_doc(course=course, subject=subject, user=professor)

        client.force_login(student)
        hub = client.get(reverse("analytics_student:materials_hub"))
        hub_content = hub.content.decode()
        assert hub.status_code == 200
        assert subject.code in hub_content
        assert "Learning materials" in hub_content
        # Table layout matching faculty course subjects, with Open action
        assert 'id="materials-subjects-table"' in hub_content
        assert ">Open</a>" in hub_content
        assert "materials-subject-grid" not in hub_content

        listing = client.get(
            reverse(
                "analytics_student:materials_list",
                kwargs={"subject_pk": subject.pk},
            )
        )
        content = listing.content.decode()
        assert listing.status_code == 200
        assert "Download" in content
        assert doc.display_title in content
        assert "Upload material" not in content
        assert "Archive" not in content

        download = client.get(
            reverse(
                "analytics_student:materials_download",
                kwargs={"subject_pk": subject.pk, "pk": doc.pk},
            )
        )
        assert download.status_code == 200
        assert download["Content-Disposition"].startswith("attachment;")


@pytest.mark.django_db
class TestIngestFullFileDownload:
    def test_keeps_full_file_ready_when_text_extraction_fails(self, professor):
        upload = SimpleUploadedFile(
            "module.pdf",
            b"%PDF-1.4 fake binary that will not extract",
            content_type="application/pdf",
        )
        with patch(
            "apps.ai.ingest.extract_pages_from_upload",
            side_effect=SourceMaterialError("Could not extract enough text"),
        ):
            doc = ingest_learning_upload(
                uploaded_file=upload,
                course_id=1,
                user=professor,
            )
        assert doc.status == LearningDocument.Status.READY
        assert doc.file
        assert doc.file.read() == b"%PDF-1.4 fake binary that will not extract"
        assert "Full file saved for download" in doc.error_message
